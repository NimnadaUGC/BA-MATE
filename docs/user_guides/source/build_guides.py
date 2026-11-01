#!/usr/bin/env python3
"""Publish app guides using the retained SKILL guide's A4 design system.

Run with a Python environment containing python-docx and lxml. The default
template is in the sibling BA-MATE-SKILL checkout; --template can override it.
PDF publishing uses the document skill renderer, independently of this script.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile
from lxml import etree

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
DEFAULT_TEMPLATE = OUT.parents[2] / 'BA-MATE-SKILL/docs/user_guide/BA-Mate-01-Getting-Started-v1.0.5.docx'
VERSION = '1.0.0-beta.3'
WIDTH = 6.614
GUIDES = [
    ('01-installation-and-ai-setup', 'Installation and AI Setup', 7),
    ('02-complete-workflow', 'Complete Workflow', 11),
    ('03-beta-testing-and-recovery', 'Beta Testing and Recovery', 6),
]


def xml(tag, attrs=None):
    node = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        node.set(qn(k), str(v))
    return node


def hyperlink(p, label, target):
    h = xml('w:hyperlink', {'r:id': p.part.relate_to(target, RT.HYPERLINK, is_external=True)})
    r = xml('w:r'); rp = xml('w:rPr')
    rp.append(xml('w:color', {'w:val': '245E7F'}))
    rp.append(xml('w:u', {'w:val': 'single'})); r.append(rp)
    t = xml('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def inline(p, text):
    for item in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))', text):
        if item.startswith('**'):
            p.add_run(item[2:-2]).bold = True
        elif item.startswith('`'):
            r = p.add_run(item[1:-1]); r.font.name = 'DejaVu Sans Mono'; r.font.size = Pt(9.5)
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', item):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', item)
            hyperlink(p, m[1], m[2])
        elif item:
            for fragment in re.split(r'((?<!Guide 1 )(?<!Guide 2 )(?<!Guide 3 )page \d+)', item):
                m = re.fullmatch(r'page (\d+)', fragment)
                if m:
                    h = xml('w:hyperlink', {'w:anchor': 'guide-page-' + m[1]})
                    r = xml('w:r'); t = xml('w:t'); t.text = fragment; r.append(t); h.append(r); p._p.append(h)
                elif fragment:
                    p.add_run(fragment)


def field(p, name, cached):
    f = xml('w:fldSimple', {'w:instr': name})
    r = xml('w:r'); t = xml('w:t'); t.text = str(cached); r.append(t); f.append(r); p._p.append(f)


def step_numbering(doc):
    root = doc.part.numbering_part.element
    aid = max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    nid = max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))] + [0]) + 1
    a = xml('w:abstractNum', {'w:abstractNumId': aid})
    a.append(xml('w:multiLevelType', {'w:val': 'singleLevel'}))
    lev = xml('w:lvl', {'w:ilvl': 0})
    for tag, value in [('start', 1), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
        lev.append(xml('w:' + tag, {'w:val': value}))
    ppr = xml('w:pPr'); ppr.append(xml('w:ind', {'w:left': 350, 'w:hanging': 350})); lev.append(ppr)
    a.append(lev)
    # OOXML orders abstract definitions before concrete number instances.
    first_num = root.find(qn('w:num'))
    if first_num is not None: root.insert(list(root).index(first_num), a)
    else: root.append(a)
    n = xml('w:num', {'w:numId': nid}); n.append(xml('w:abstractNumId', {'w:val': aid}))
    cleanup_id = root.find(qn('w:numIdMacAtCleanup'))
    if cleanup_id is not None: root.insert(list(root).index(cleanup_id), n)
    else: root.append(n)
    return nid


def table(doc, lines):
    rows = [[c.strip() for c in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    ratios = [.29, .71]
    if rows[0][0] in ('Sidebar location', 'Location or record', 'Practice check'):
        ratios = [.40, .60]
    if rows[0][0] == 'Symptom':
        ratios = [.31, .69]
    t = doc.add_table(rows=1, cols=len(rows[0])); t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, ratio in zip(t.columns, ratios): c.width = Inches(WIDTH * ratio)
    borders = xml('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(xml('w:' + side, {'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
    t._tbl.tblPr.append(borders)
    for i, vals in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        pr = row._tr.get_or_add_trPr(); pr.append(xml('w:cantSplit'))
        if i == 0: pr.append(xml('w:tblHeader'))
        for j, val in enumerate(vals):
            c = row.cells[j]; c.width = Inches(WIDTH * ratios[j]); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcp = c._tc.get_or_add_tcPr(); mar = xml('w:tcMar')
            for side, size in [('top', 110), ('bottom', 110), ('left', 135), ('right', 135)]:
                mar.append(xml('w:' + side, {'w:w': size, 'w:type': 'dxa'}))
            tcp.append(mar); tcp.append(xml('w:shd', {'w:fill': '193F5F' if i == 0 else ('F3F6F7' if i % 2 == 0 else 'FFFFFF')}))
            p = c.paragraphs[0]; p.style = 'Table Text'; inline(p, val)
            if i == 0:
                for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph(style='Table Gap')


def chrome(doc, number, short_title, pages):
    s = doc.sections[0]
    for hdr in (s.header, s.first_page_header):
        for child in list(hdr._element): hdr._element.remove(child)
        p = hdr.add_paragraph(); p.style = 'Header'
        if hdr == s.header: p.add_run(f'BA Mate   /   App Guide {number}   /   {short_title}')
    # The property accessor returns separate wrappers, so identify the ordinary header explicitly.
    s.header.paragraphs[0].text = f'BA Mate   /   App Guide {number}   /   {short_title}'
    for ftr in (s.footer, s.first_page_footer):
        for child in list(ftr._element): ftr._element.remove(child)
        p = ftr.add_paragraph(); p.style = 'Footer'
        p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('Application beta.3  ·  1 November 2026\t')
        field(p, 'PAGE', 1); p.add_run(' / '); field(p, 'NUMPAGES', pages)
    cp = doc.core_properties
    cp.title = 'BA Mate ' + short_title; cp.subject = 'Application beta tester guide'
    cp.author = cp.last_modified_by = 'BA Mate research project'; cp.language = 'en-GB'
    cp.keywords = f'BA-MATE, {VERSION}, application, beta, user guide'
    cp.created = cp.modified = datetime(2026, 11, 1, tzinfo=timezone.utc)
    for old in list(doc.settings.element.findall(qn('w:updateFields'))): doc.settings.element.remove(old)
    doc.settings.element.append(xml('w:updateFields', {'w:val': 'true'}))


def build(template, slug, title, pages, number):
    doc = Document(template)
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'): doc._element.body.remove(child)
    chrome(doc, number, title, pages)
    p = doc.add_paragraph(style='Figure'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(OUT / 'assets/image1.png'), width=Inches(1.1))
    shape._inline.docPr.set('descr', 'BA Mate logo')
    lines = (HERE / (slug + '.md')).read_text().splitlines()
    num_id = None; i = 0; logical_page = 1; pending_break = False
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line == '<!-- page -->':
            logical_page += 1; pending_break = True; num_id = None; continue
        if line.startswith('```'):
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); p.paragraph_format.keep_together = True
            for j, txt in enumerate(code):
                r = p.add_run(('\n' if j else '') + txt); r.font.name = 'DejaVu Sans Mono'; r.font.size = Pt(9.5)
            continue
        if line.startswith('|'):
            group = [line]
            while i < len(lines) and lines[i].startswith('|'):
                group.append(lines[i].strip()); i += 1
            table(doc, group); num_id = None; continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title'); continue
        if line.startswith('## '):
            p = doc.add_paragraph(line[3:], 'Heading 1')
            if pending_break:
                p.paragraph_format.page_break_before = True
                start = xml('w:bookmarkStart', {'w:id': 1000 + logical_page, 'w:name': 'guide-page-' + str(logical_page)})
                p._p.insert(1, start); p._p.append(xml('w:bookmarkEnd', {'w:id': 1000 + logical_page}))
                pending_break = False
            num_id = None; continue
        if line.startswith('Application guide '):
            doc.add_paragraph(line, 'Guide Meta'); continue
        m = re.match(r'^(\d+)\. (.*)', line)
        if m:
            if m[1] == '1' or num_id is None: num_id = step_numbering(doc)
            p = doc.add_paragraph(); inline(p, m[2])
            np = xml('w:numPr'); np.append(xml('w:ilvl', {'w:val': 0})); np.append(xml('w:numId', {'w:val': num_id})); p._p.get_or_add_pPr().append(np)
            p.paragraph_format.space_after = Pt(7)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); inline(p, line[2:]); num_id = None
        else:
            p = doc.add_paragraph(); inline(p, line); num_id = None
    path = OUT / f'BA-Mate-App-{slug}-v{VERSION}.docx'
    doc.save(path)
    # Preserve opaque source parts exactly. Body/chrome/properties/numbering are intended edit slots.
    with zipfile.ZipFile(template) as z: baseline = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(path) as z: data = {n: z.read(n) for n in z.namelist()}
    preserve = [n for n in baseline if n.startswith(('word/styles', 'word/theme/', 'word/fontTable', 'word/_rels/fontTable', 'word/fonts/', 'word/media/'))]
    for n in preserve: data[n] = baseline[n]
    app = etree.fromstring(data['docProps/app.xml'])
    ns = {'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}
    page_count = app.find('ep:Pages', ns)
    if page_count is not None: page_count.text = str(pages)
    data['docProps/app.xml'] = etree.tostring(app, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, value in data.items(): z.writestr(n, value)
    with zipfile.ZipFile(path) as z:
        assert all(z.read(n) == baseline[n] for n in preserve)
    return {'file': path.name, 'source': slug + '.md', 'expected_pages': pages, 'preserved_parts': len(preserve), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--template', type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    if not args.template.exists(): raise SystemExit('Supply --template pointing to the retained SKILL getting started DOCX.')
    records = [build(args.template, *g, i) for i, g in enumerate(GUIDES, 1)]
    manifest = {'application': VERSION, 'framework': 'F4-draft.1 (v1.0.0)', 'skill_companion': 'v1.0.5',
                'guide_date': '2026-11-01', 'template': args.template.name,
                'template_sha256': hashlib.sha256(args.template.read_bytes()).hexdigest(), 'docx': records}
    (HERE / 'build-manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps(records, indent=2))
