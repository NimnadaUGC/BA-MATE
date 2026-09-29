import base64
import hashlib
import io
import zipfile
from pydantic import BaseModel, Field
from fastapi import HTTPException

MAX_BYTES = 25 * 1024 * 1024
MAX_TEXT = 200000
MEDIA_SUFFIXES = {
    'png': 'image', 'jpg': 'image', 'jpeg': 'image', 'gif': 'image', 'webp': 'image', 'bmp': 'image', 'tif': 'image', 'tiff': 'image',
    'mp3': 'audio', 'wav': 'audio', 'm4a': 'audio', 'aac': 'audio', 'ogg': 'audio', 'flac': 'audio',
    'mp4': 'video', 'mov': 'video', 'm4v': 'video', 'webm': 'video', 'mkv': 'video',
}

class SourceUpload(BaseModel):
    name: str = Field(min_length=1, max_length=255)
    data: str = Field(max_length=35 * 1024 * 1024)


def extract_source(request: SourceUpload):
    try:
        data = base64.b64decode(request.data, validate=True)
        if not data or len(data) > MAX_BYTES:
            raise ValueError('Use a non-empty file of at most 25 MB.')
        suffix = request.name.lower().rsplit('.', 1)[-1]
        passages = []
        limitations = ['Extraction may omit layout, images, headers or text boxes. Review against the original.']
        media_kind = 'document'
        if suffix == 'pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                limitations = ['The original PDF was preserved, but it is password protected and cannot be used as searchable evidence.']
            elif len(reader.pages) > 250:
                limitations = ['The original PDF was preserved, but PDFs longer than 250 pages are not extracted. Split it for searchable evidence.']
            else:
                for index, page in enumerate(reader.pages):
                    content = page.extract_text() or ''
                    if content.strip():
                        passages.append({'locator': f'Page {index + 1}', 'text': content.strip()})
                if not passages:
                    limitations = ['The original PDF was preserved, but no text was found. Add a transcription or OCR copy before using it as evidence.']
        elif suffix == 'docx':
            from docx import Document
            archive = zipfile.ZipFile(io.BytesIO(data))
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('The expanded document is too large.')
            document = Document(io.BytesIO(data))
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    passages.append({'locator': f'Paragraph {index + 1}', 'text': paragraph.text})
            for index, table in enumerate(document.tables):
                passages.append({'locator': f'Table {index + 1}', 'text': '\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)})
        elif suffix in {'txt', 'md', 'csv', 'json'}:
            content = data.decode('utf-8-sig')
            for index, paragraph in enumerate(content.split('\n\n')):
                if paragraph.strip():
                    passages.append({'locator': f'Block {index + 1}', 'text': paragraph.strip()})
        elif suffix in {'xlsx', 'xls'}:
            if suffix == 'xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                try:
                    for sheet in workbook.worksheets:
                        rows = []
                        for row in sheet.iter_rows(values_only=True):
                            values = ['' if value is None else str(value) for value in row]
                            if any(value.strip() for value in values):
                                rows.append(' | '.join(values))
                        if rows:
                            passages.append({'locator': f'Sheet {sheet.title}', 'text': '\n'.join(rows)})
                finally:
                    workbook.close()
            else:
                import xlrd
                workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
                try:
                    for sheet in workbook.sheets():
                        rows = []
                        for row_index in range(sheet.nrows):
                            values = [str(sheet.cell_value(row_index, column)).strip() for column in range(sheet.ncols)]
                            if any(values):
                                rows.append(' | '.join(values))
                        if rows:
                            passages.append({'locator': f'Sheet {sheet.name}', 'text': '\n'.join(rows)})
                finally:
                    workbook.release_resources()
            media_kind = 'spreadsheet'
            if not passages:
                limitations = ['The original spreadsheet was preserved, but it contains no readable cell values.']
        elif suffix in MEDIA_SUFFIXES:
            media_kind = MEDIA_SUFFIXES[suffix]
            label = {'image': 'OCR', 'audio': 'a transcript', 'video': 'a transcript'}[media_kind]
            limitations = [f'The original {media_kind} file was preserved. Add {label} before using it as searchable evidence.']
        else:
            raise ValueError('Supported files: PDF, DOCX, TXT, Markdown, CSV, JSON, Excel, common images, audio and video.')
        text = '\n\n'.join(f"[{item['locator']}] {item['text']}" for item in passages)
        if len(text) > MAX_TEXT:
            raise ValueError('Extracted text exceeds 200,000 characters. Split this source before importing.')
        return {'name': request.name, 'content': text, 'passages': passages, 'sha256': hashlib.sha256(data).hexdigest(), 'byte_size': len(data), 'original_base64': request.data,
                'status': 'Needs review', 'approved': False,
                'media_kind': media_kind, 'extractable': bool(text.strip()), 'limitations': limitations}
    except (ValueError, UnicodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(422, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(422, 'This document could not be extracted. Try exporting it as text.') from exc
