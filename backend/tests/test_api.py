import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
from fastapi.testclient import TestClient
from app.main import app
from app import runtime
from app.contracts import RunRequest
from app.workflow import prepare_messages, schema_for_provider


def proposal(**changes):
    return {'summary': 'Review this proposal', 'questions': [], 'requirements': [], 'stories': [], 'findings': [], 'diagrams': [], 'sections': [], 'limitations': [], **changes}

class RealWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path_patch = patch.object(runtime, 'DATA_DIR', Path(self.temp.name))
        self.path_patch.start()
        self.config_patch = patch.object(runtime, 'config', runtime.ProviderConfig())
        self.config_patch.start()
        self.client = TestClient(app)
        self.client.headers['x-ba-session'] = runtime.SESSION_TOKEN
        self.request = {'request_id': 'test-run', 'project_id': 'project-A', 'task': 'clarify', 'prompt': 'Find missing exception rules',
                        'context': [{'id': 'SRC-1', 'type': 'Source', 'content': '[Page 2] Managers approve leave requests. Escalation rules are not yet decided.', 'version': 1, 'provenance': 'Workshop notes', 'priority': 10}], 'cloud_allowed': False}

    def tearDown(self):
        self.config_patch.stop(); self.path_patch.stop(); self.temp.cleanup()

    def test_source_extraction_retains_original_and_requires_review(self):
        content = 'Customer submits an order.\n\nA manager reviews refunds.'
        encoded = base64.b64encode(content.encode()).decode()
        response = self.client.post('/api/sources/extract', json={'name': 'notes.txt', 'data': encoded})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertIn(content.split('\n')[0], body['content'])
        self.assertEqual(body['original_base64'], encoded)
        self.assertFalse(body['approved'])
        self.assertEqual(len(body['sha256']), 64)
        self.assertEqual(body['passages'][1]['locator'], 'Block 2')

    def test_docx_paragraphs_and_tables_are_extracted(self):
        from docx import Document
        doc = Document(); doc.add_paragraph('Business objective: reduce manual re-entry.')
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Actor'; table.cell(0, 1).text = 'Analyst'
        buffer = io.BytesIO(); doc.save(buffer)
        response = self.client.post('/api/sources/extract', json={'name': 'brief.docx', 'data': base64.b64encode(buffer.getvalue()).decode()})
        self.assertEqual(response.status_code, 200)
        self.assertIn('Actor | Analyst', response.json()['content'])

    def test_images_are_not_misrepresented_as_processed_text(self):
        response = self.client.post('/api/sources/extract', json={'name': 'scan.png', 'data': base64.b64encode(b'fake').decode()})
        self.assertEqual(response.status_code, 422)

    def test_workspace_conflicts_do_not_overwrite_saved_work(self):
        first = {'schemaVersion': 6, 'projects': [{'id': 'A', 'name': 'Actual project'}]}
        self.assertEqual(self.client.put('/api/workspace', json={'expected_revision': 0, 'state': first}).status_code, 200)
        conflict = self.client.put('/api/workspace', json={'expected_revision': 0, 'state': {'schemaVersion': 6, 'projects': []}})
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(self.client.get('/api/workspace').json()['state'], first)

    def test_empty_workspace_is_a_valid_saved_state(self):
        state = {'schemaVersion': 6, 'projects': []}
        self.client.put('/api/workspace', json={'expected_revision': 0, 'state': state})
        self.assertEqual(self.client.get('/api/workspace').json()['state'], state)

    def test_session_and_origin_boundary(self):
        response = TestClient(app).get('/api/workspace')
        self.assertEqual(response.status_code, 401)
        self.assertEqual(self.client.get('/api/session', headers={'origin': 'https://unrelated.example'}).status_code, 403)
        self.assertEqual(self.client.get('/api/session', headers={'host': 'attacker.example'}).status_code, 403)

    def test_provider_keys_do_not_enter_disk_settings(self):
        for provider in ['openai', 'gemini', 'openrouter', 'huggingface']:
            with self.subTest(provider=provider):
                secret = f'secret-{provider}-value'
                response = self.client.put('/api/settings', json={'provider': provider, 'model': 'test-model', 'api_key': secret})
                self.assertEqual(response.status_code, 200)
                self.assertNotIn(secret, response.text)
                self.assertNotIn(secret, (Path(self.temp.name) / 'provider.json').read_text())

    def test_session_credential_can_be_explicitly_removed(self):
        self.client.put('/api/settings', json={'provider': 'openrouter', 'model': 'test/model', 'api_key': 'secret-test-value'})
        response = self.client.delete('/api/settings/credential')
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['has_key'])
        self.assertNotIn('secret-test-value', (Path(self.temp.name) / 'provider.json').read_text())

    def test_full_evidence_and_governance_reach_provider(self):
        result = proposal(questions=[{'question': 'Who handles overdue requests?', 'rationale': 'Escalation is undefined', 'source_ids': ['SRC-1'], 'priority': 'Blocking'}])
        fake = AsyncMock(return_value={'text': json.dumps(result), 'model': 'test-model', 'provider': 'ollama', 'usage': {'input_tokens': 100, 'output_tokens': 70, 'cost': None}})
        with patch('app.workflow.ModelProvider.complete', fake):
            response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 200, response.text)
        messages = fake.call_args.args[0]
        self.assertIn('Managers approve leave requests', json.dumps(messages))
        self.assertIn('untrusted data', messages[0]['content'])
        self.assertEqual(response.json()['result']['context_receipt']['included_reference_ids'], ['SRC-1'])
        self.assertNotIn('confidence', response.json()['metadata'])
        self.assertIsNone(self.client.get('/api/workspace').json()['state'])

    def test_invented_evidence_reference_is_rejected(self):
        result = proposal(questions=[{'question': 'What deadline?', 'rationale': 'Missing', 'source_ids': ['NOT-SUPPLIED'], 'priority': 'Blocking'}])
        with patch('app.workflow.ModelProvider.complete', AsyncMock(return_value={'text': json.dumps(result), 'model': 'test', 'provider': 'ollama', 'usage': {}})):
            response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 502)
        self.assertIn('not supplied', response.text)

    def test_malformed_model_output_is_not_applied(self):
        with patch('app.workflow.ModelProvider.complete', AsyncMock(return_value={'text': 'not JSON', 'model': 'test', 'provider': 'ollama', 'usage': {}})):
            response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 502)

    def test_cloud_processing_requires_project_preference(self):
        runtime.config = runtime.ProviderConfig(provider='openrouter', model='test')
        response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 409)

    def test_confidential_context_never_uses_cloud_even_when_allowed(self):
        runtime.config = runtime.ProviderConfig(provider='huggingface', model='test')
        self.request['cloud_allowed'] = True
        self.request['context'][0]['confidential'] = True
        self.assertEqual(self.client.post('/api/v1/workflow/run', json=self.request).status_code, 409)

    def test_receipt_discloses_partial_evidence(self):
        self.request['context'][0]['content'] = 'Long business evidence. ' * 7000
        messages, receipt = prepare_messages(RunRequest.model_validate(self.request), runtime.ProviderConfig(), schema_for_provider())
        self.assertEqual(receipt['partial_reference_ids'], ['SRC-1'])
        self.assertLess(receipt['estimated_tokens'] + runtime.config.output_tokens, runtime.config.context_tokens)

    def test_generic_condition_does_not_receive_framework_prompt(self):
        self.request['condition'] = 'Generic AI'
        messages, receipt = prepare_messages(RunRequest.model_validate(self.request), runtime.ProviderConfig(), schema_for_provider())
        self.assertNotIn('evidence-grounded workflow', messages[0]['content'])
        self.assertFalse(receipt['governance_included'])

    def test_changed_evaluation_configuration_is_rejected_before_inference(self):
        self.request['expected_configuration'] = 'old-settings'
        with patch('app.workflow.ModelProvider.complete', AsyncMock()) as provider:
            response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 409)
        provider.assert_not_called()

    def test_missing_confirmed_quantity_is_visible_for_human_review(self):
        self.request['task'] = 'requirements'
        self.request['context'].append({'id': 'CQ-1', 'type': 'Clarification', 'content': 'Status: Answered\nAnswer: Reservations expire after 48 hours.'})
        result = proposal(requirements=[{'title': 'Expiry', 'statement': 'Reservations expire.', 'kind': 'Functional', 'rationale': 'Policy', 'acceptance_criteria': ['Reservation expires.'], 'source_ids': ['CQ-1'], 'assumptions': []}])
        with patch('app.workflow.ModelProvider.complete', AsyncMock(return_value={'text': json.dumps(result), 'model': 'test', 'provider': 'ollama', 'usage': {}})):
            response = self.client.post('/api/v1/workflow/run', json=self.request)
        self.assertEqual(response.status_code, 200)
        self.assertIn('48 hours', ' '.join(response.json()['result']['proposal']['limitations']))

    def test_retired_endpoints_cannot_return_mock_results(self):
        self.assertEqual(self.client.post('/api/v1/generation/artifacts', json={}).status_code, 410)

if __name__ == '__main__':
    unittest.main()
